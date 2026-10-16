from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).parent / 'SavEx_Quantum_MFB_Partnership_Questionnaire.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.1
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'SavEx Quantum MFB Partnership Questionnaire'
doc.core_properties.subject = 'Proposed customer accounts payments reconciliation and settlement arrangement'
doc.core_properties.author = 'SavEx'
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('SavEx and Quantum MFB  |  Page ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def question(n, title, detail):
    doc.add_paragraph(f'{n}  {title}', 'Heading 2')
    p = doc.add_paragraph(detail)
    p.paragraph_format.keep_with_next = True
    for i in range(2):
        p = doc.add_paragraph('Response: ' if i == 0 else '')
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = i == 0
        p.paragraph_format.line_spacing = 1.0
        borders = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for k, v in [('val', 'single'), ('sz', '3'), ('color', 'C8C8C8'), ('space', '4')]:
            bottom.set(qn('w:'+k), v)
        borders.append(bottom); p._p.get_or_add_pPr().append(borders)

title = doc.add_paragraph('SavEx Quantum MFB\nPartnership Questionnaire', 'Title')
for border in list(title._p.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.add_paragraph('Please confirm the proposed banking and operational arrangements below so we can prepare the definitive agreement and integration plan. Provide supporting documents where available. Mark unresolved items as pending and identify the responsible contact.')
doc.add_paragraph('For discussion only. Completing this questionnaire does not authorise launch, movement of customer funds or create a binding partnership. All services remain subject to written agreements and required approvals.')
doc.add_paragraph('Bank legal name: ______________________________________________\nSavEx operating company: ________________________________________\nRespondent and role: _____________________  Date: _________________\nEmail and telephone: ____________________________________________')
doc.add_paragraph('Customer accounts and funds', 'Heading 1')
question(1, 'Approved services', 'Which customer accounts, funding, withdrawals, savings and savings-circle services can Quantum provide? State relevant licence scope, geographic limits and approvals still required.')
question(2, 'Account ownership', 'Will customers hold individual accounts or participate in a pooled arrangement? Identify the legal holder, beneficial ownership and who contracts with each customer.')
question(3, 'Account structure', 'Which accounts will receive customer funds, support withdrawals, receive settlements and collect fees? Identify permitted uses and any correspondent-bank relationship.')
question(4, 'Protection of customer funds', 'How will customer funds be separated from operating funds? Explain applicable deposit-protection treatment, conditions and exclusions.')

doc.add_page_break()
doc.add_paragraph('Payments reconciliation and settlement', 'Heading 1')
doc.add_paragraph('Please describe the full funds flow and attach a diagram or sample reports where possible.')
question(5, 'Funding and withdrawal flow', 'How will funds move between customer accounts, SavEx, Quantum and any payment provider? Who supplies payout liquidity and what prefunding or minimum balances apply?')
question(6, 'Payment confirmation', 'Which record proves a payment is final? Define queued, pending, successful, failed, reversed and restricted outcomes, including recipient-credit tracing.')
question(7, 'Timeouts and corrections', 'How are timeouts, duplicates, rejected transfers and reversals handled? Who may release reservations or approve adjustments, and what maker-checker controls apply?')
question(8, 'Settlement timetable', 'Specify destination accounts, settlement cycles, cut-off times, business-day rules, gross or net treatment, fees, taxes, reserves and delayed-settlement escalation.')
question(9, 'Reconciliation evidence', 'Which APIs, statements and transaction or settlement reports will Quantum supply? Confirm shared references, amount, currency, fee, status and date fields, delivery frequency and availability.')
question(10, 'Exception ownership', 'Who investigates unmatched payments, balance differences and missing bank receipts? Propose severity levels, response and resolution targets, and evidence required to close a case.')

doc.add_page_break()
doc.add_paragraph('Responsibilities integration and launch', 'Heading 1')
question(11, 'Customer verification and compliance', 'Allocate onboarding, KYC approval, account tiers, transaction limits, AML monitoring, sanctions checks, fraud escalation and regulatory reporting between the parties.')
question(12, 'Technical access and security', 'What sandbox and production APIs, authentication, fixed-IP requirements, notifications and statements are available? Identify security, recovery and acceptance tests required before launch.')
question(13, 'Data and customer support', 'Define data-controller and processor roles, lawful sharing, retention, incident escalation and any cross-border processing. Who handles customer complaints and bank investigations?')
question(14, 'Commercial terms and losses', 'Propose setup and recurring charges, transaction fees, revenue sharing and taxes. How should fraud, processing errors, outages and third-party losses be allocated?')
question(15, 'Existing funds and providers', 'Can Flutterwave or other providers remain in the arrangement? Specify how existing balances and restricted withdrawals must be reconciled before any migration, without duplicate credits or payouts.')
question(16, 'Launch and exit conditions', 'List legal, regulatory, commercial and technical sign-offs for a controlled pilot. Define ongoing audit access, termination, customer-fund return or transfer, final reconciliation and record retention.')
doc.add_paragraph('Supporting documents and next meeting', 'Heading 2')
doc.add_paragraph('Please attach the proposed account and funds-flow structure, API documentation, sample statements, pricing, service levels, compliance requirements and data-processing terms.')
doc.add_paragraph('Outstanding items and owners: _____________________________________\n______________________________________________________________\nProposed follow-up date: __________________________________________')
doc.save(out)
print(out)
